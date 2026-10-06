#!/usr/bin/env python3
"""Gera areas-data.js e img/areas/*.png a partir de dados/Cursos_Etecs_Zona_Leste.xlsx.

Uso (na raiz do projeto):  pip install openpyxl && python3 tools/xlsx_para_areas.py
Falha com mensagem clara se a planilha sair do formato esperado.
"""
import json, re, sys, unicodedata, zipfile
import openpyxl

XLSX = "dados/Cursos_Etecs_Zona_Leste.xlsx"
ADHEMAR = "Etec Professor Adhemar Batista Heméritas"
# Aba "Fontes e observações": "Quando nenhuma Etec da zona leste tem o curso exato da área
# (Licenciaturas, Agropecuária/Zootecnia e Estética), foram indicados os cursos mais próximos."
SEM_CURSO_EXATO = {"educacao-e-licenciaturas", "agronegocio-e-zootecnia", "estetica-beleza-e-bem-estar"}


def slug(s):
    s = unicodedata.normalize("NFD", s.lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def erro(msg):
    sys.exit(f"ERRO: {msg}")


wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["Cursos"]
wd = wb["Distâncias e fontes"]

# ---- ids das áreas do teste (quiz-data.js) ----
ids_teste = re.findall(r'"id":\s*"([^"]+)"', open("js/quiz-data.js", encoding="utf-8").read())

# ---- unidades (aba "Distâncias e fontes", linhas 11+) ----
unidades = {}
for r in range(11, wd.max_row + 1):
    nome = wd.cell(r, 1).value
    if not nome:
        continue
    unidades[nome.strip()] = {
        "id": slug(nome),
        "nome": nome.strip(),
        "endereco": wd.cell(r, 2).value.strip(),
        "distanciaKm": float(wd.cell(r, 3).value),
        "lat": float(wd.cell(r, 4).value),
        "lng": float(wd.cell(r, 5).value),
        "precisao": wd.cell(r, 6).value.strip(),
        "url": (wd.cell(r, 7).value or "").strip(),
    }
if ADHEMAR not in unidades:
    erro("Adhemar não encontrada na aba de distâncias")

TAG = re.compile(r"^(.*?)\s*\((EaD|M-Tec-N|M-Tec)(?:,[^)]*)?\)\s*$")
MOD = {"EaD": "ead", "M-Tec": "mtec", "M-Tec-N": "mtecn"}
SUFIXO = re.compile(r"\s+–\s+(opç(?:ão|ões) relacionad[ao]s?.*)$")


def parse_unidade(cell, linha):
    linhas = [l.strip() for l in str(cell).split("\n") if l.strip()]
    if len(linhas) != 4 or not linhas[2].startswith("Curso: ") or not linhas[3].startswith("Distância aprox.:"):
        erro(f"formato inesperado na célula da linha {linha}: {linhas}")
    nome, endereco, curso, dist = linhas
    if nome not in unidades:
        erro(f"unidade '{nome}' (linha {linha}) não está na aba de distâncias")
    if unidades[nome]["endereco"] != endereco:
        erro(f"endereço divergente para '{nome}' na linha {linha}")
    corpo = curso[len("Curso: "):]
    relacionada, nota = False, ""
    m = SUFIXO.search(corpo)
    if m:
        relacionada, nota, corpo = True, m.group(1), corpo[: m.start()]
    cursos = []
    for item in [c.strip() for c in corpo.split(";") if c.strip()]:
        mt = TAG.match(item)
        cursos.append({"nome": (mt.group(1) if mt else item).strip(),
                       "modalidade": MOD[mt.group(2)] if mt else None,
                       "relacionada": relacionada})
    km = float(re.search(r"([\d,]+)\s*km", dist).group(1).replace(",", "."))
    if abs(km - unidades[nome]["distanciaKm"]) > 0.06:
        erro(f"distância divergente para '{nome}' linha {linha}: {km} x {unidades[nome]['distanciaKm']}")
    return {"unidade": unidades[nome]["id"], "cursos": cursos, "notaCursos": nota}


areas = {}
por_id = {u["id"]: u for u in unidades.values()}
for r in range(3, ws.max_row + 1):
    nome = ws.cell(r, 1).value
    if not nome:
        continue
    aid = slug(nome)
    if aid not in ids_teste:
        erro(f"área '{nome}' (linha {r}) não corresponde a nenhum id de quiz-data.js")
    atuacao = [re.sub(r"^\d+\.\s*", "", l.strip()) for l in str(ws.cell(r, 4).value).split("\n") if l.strip()]
    if len(atuacao) != 4:
        erro(f"{nome}: esperava 4 áreas de atuação, achei {len(atuacao)}")
    ofertas = [parse_unidade(ws.cell(r, c).value, r) for c in (6, 7, 8)]
    ordenadas = sorted(ofertas, key=lambda o: (o["unidade"] != slug(ADHEMAR), por_id[o["unidade"]]["distanciaKm"], por_id[o["unidade"]]["nome"]))
    if [o["unidade"] for o in ordenadas] != [o["unidade"] for o in ofertas]:
        erro(f"{nome}: a ordem da planilha difere da regra (Adhemar primeiro, depois distância)")
    if aid in SEM_CURSO_EXATO:
        for o in ordenadas:
            for c in o["cursos"]:
                c["relacionada"] = True
    areas[aid] = {
        "id": aid,
        "semCursoExato": aid in SEM_CURSO_EXATO,
        "nome": nome.strip(),
        "explicacao": ws.cell(r, 2).value.strip(),
        "atuacao": atuacao,
        "salario": [l.strip() for l in str(ws.cell(r, 5).value).split("\n") if l.strip()],
        "observacao": (ws.cell(r, 9).value or "").strip(),
        "icone": f"../img/areas/{aid}.png",
        "ofertas": ordenadas,
    }
if sorted(areas) != sorted(ids_teste):
    erro(f"áreas ausentes: {set(ids_teste) - set(areas)}")

avisos = {
    "ofertaMuda": "A oferta de cursos muda a cada semestre. Antes de se inscrever, confirme no site oficial: vestibulinho.etec.sp.gov.br",
    "salario": "Faixas salariais são estimativas aproximadas para a cidade de São Paulo. Variam conforme cargo, empresa, experiência e formação técnica ou superior.",
    "distancia": "Distância aproximada em linha reta a partir da Etec Professor Adhemar Batista Heméritas. Não representa percurso, tempo de viagem ou transporte público.",
    "selecao": "Seleção de unidades da zona leste com cursos pertinentes; não é um levantamento exaustivo de todas as ETECs. Cadastro de cursos não garante abertura de vagas.",
    "modalidades": "M-Tec = Ensino Médio integrado ao técnico; M-Tec-N = integrado no período noturno; EaD = online.",
    "consulta": "Dados consultados em 06/10/2026.",
}

with open("js/areas-data.js", "w", encoding="utf-8") as f:
    f.write("// Gerado por tools/xlsx_para_areas.py a partir de dados/Cursos_Etecs_Zona_Leste.xlsx. Não edite à mão.\n")
    f.write("const ADHEMAR_ID = " + json.dumps(slug(ADHEMAR)) + ";\n")
    f.write("const UNIDADES = " + json.dumps(por_id, ensure_ascii=False, indent=1) + ";\n")
    f.write("const AREAS_INFO = " + json.dumps(areas, ensure_ascii=False, indent=1) + ";\n")
    f.write("const AVISOS = " + json.dumps(avisos, ensure_ascii=False, indent=1) + ";\n")

# ---- ícones: a imagem N fica ancorada na linha N+2 da aba Cursos ----
z = zipfile.ZipFile(XLSX)
d = z.read("xl/drawings/drawing1.xml").decode("utf-8")
rels = {i: t.split("/")[-1] for t, i in re.findall(r'Target="([^"]+)" Id="(rId\d+)"', z.read("xl/drawings/_rels/drawing1.xml.rels").decode())}
linha_para_id = {r: slug(ws.cell(r, 1).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value}
n = 0
for a in re.findall(r"<oneCellAnchor>.*?</oneCellAnchor>", d, flags=re.S):
    linha = int(re.search(r"<row>(\d+)</row>", a).group(1)) + 1
    img = rels[re.search(r'r:embed="(rId\d+)"', a).group(1)]
    if linha in linha_para_id:
        open(f"img/areas/{linha_para_id[linha]}.png", "wb").write(z.read("xl/media/" + img))
        n += 1
if n != len(areas):
    erro(f"esperava {len(areas)} ícones, extraí {n}")
print(f"OK: {len(areas)} áreas, {len(unidades)} unidades, {n} ícones")
